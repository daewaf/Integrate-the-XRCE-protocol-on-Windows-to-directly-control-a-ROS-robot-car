# -*- coding: utf-8 -*-
"""生成《FishBot Windows 原生控制 · 启动与关闭手册》Word 文档。"""
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

EAST_ASIA_FONT = "微软雅黑"
LINK_COLOR = "0563C1"


def set_run_font(run, size=10.5, bold=False, color=None, font_name=EAST_ASIA_FONT):
    run.font.name = font_name
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    for attr in ("w:ascii", "w:hAnsi", "w:eastAsia"):
        rfonts.set(qn(attr), font_name)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.font.bold = bold
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)


def shade(paragraph, fill="F2F2F2"):
    ppr = paragraph._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    ppr.append(shd)


def body(doc, text, bold=False, color=None, size=10.5):
    p = doc.add_paragraph()
    run = p.add_run(text)
    set_run_font(run, size=size, bold=bold, color=color)
    return p


def bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    run = p.add_run(text)
    set_run_font(run)
    return p


def code(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.18)
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(6)
    shade(p)
    for i, line in enumerate(text.splitlines()):
        if i:
            p.add_run().add_break()
        run = p.add_run(line)
        set_run_font(run, size=9, font_name="Consolas")
    return p


def heading(doc, text, level):
    h = doc.add_heading(text, level=level)
    for run in h.runs:
        set_run_font(run, size=None, bold=True)
    return h


def table(doc, headers, rows, widths=None):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Light Grid Accent 1"
    # 表头跨页重复，且行内内容不拆分
    tr_pr = t.rows[0]._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)
    for i, text in enumerate(headers):
        cell = t.rows[0].cells[i]
        cell.text = ""
        run = cell.paragraphs[0].add_run(text)
        set_run_font(run, bold=True)
    for row in rows:
        new_row = t.add_row()
        cells = new_row.cells
        cant_split = OxmlElement("w:cantSplit")
        new_row._tr.get_or_add_trPr().append(cant_split)
        for i, text in enumerate(row):
            cells[i].text = ""
            run = cells[i].paragraphs[0].add_run(str(text))
            set_run_font(run, size=9.5)
    if widths:
        for row in t.rows:
            for i, w in enumerate(widths):
                row.cells[i].width = Inches(w)
    doc.add_paragraph()
    return t


doc = Document()
section = doc.sections[0]
section.left_margin = Inches(0.9)
section.right_margin = Inches(0.9)
section.top_margin = Inches(0.8)
section.bottom_margin = Inches(0.8)

normal = doc.styles["Normal"]
normal.font.name = EAST_ASIA_FONT
normal.font.size = Pt(10.5)
normal._element.rPr.rFonts.set(qn("w:eastAsia"), EAST_ASIA_FONT)

title = doc.add_heading("FishBot Windows 原生控制", level=0)
for run in title.runs:
    set_run_font(run, bold=True)
sub = doc.add_paragraph()
run = sub.add_run("启动与关闭手册（xrce_windows_agent）")
set_run_font(run, size=12, color="595959")
sub.alignment = WD_ALIGN_PARAGRAPH.CENTER

body(doc, "适用：纯 Python XRCE-DDS agent。无需安装 ROS2，无需经过 WSL Ubuntu。")
body(doc, "更新日期：2026-10-26", color="595959", size=9.5)

heading(doc, "1. 这套工具是什么", 1)
body(doc, "Windows 上的 Python agent 直接与主控板完成 XRCE-DDS 握手、实体创建、Twist 下发，"
          "并解析主控板上报的 odom / imu 数据。WASD 遥控通过本机 TCP 控制口接入。")
code(doc, "WASD 终端 ──TCP 8898──> agent ──UDP 8888──> 车主控板\n"
          "                          <── odom / imu ────┘")
table(doc, ["组件", "作用"], [
    ("xrce/agent.py", "XRCE-DDS agent（协议转换核心）"),
    ("tools/teleop_wasd.py", "终端 WASD 遥控"),
    ("start_fishbot.py", "一键启动 agent + WASD，退出时关闭 agent"),
], widths=[2.2, 4.0])

heading(doc, "2. 前提条件", 1)
bullet(doc, "Windows 已安装 Python 3.10（本项目实测版本）。")
bullet(doc, "主控板 WiFi 与电脑在同一网络。")
bullet(doc, "主控板“上位机 IP” = 电脑当前 WLAN 地址，端口 8888。")
bullet(doc, "电脑上不要同时运行 WSL / Docker 的 micro_ros_agent（会占用 8888）。")
table(doc, ["端口", "协议", "用途"], [
    ("8888", "UDP", "XRCE-DDS（主控板 ↔ agent）"),
    ("8898", "TCP", "本机控制口（WASD ↔ agent）"),
], widths=[1.0, 1.0, 3.8])

heading(doc, "3. 一键启动（推荐）", 1)
code(doc, 'cd "C:\\Users\\kw188\\Desktop\\ros小车\\xrce_windows_agent"\n'
          "python start_fishbot.py")
body(doc, "启动流程：", bold=True)
for item in [
    "打印本机 IP（这个地址要填到主控板的“上位机 IP”）。",
    "后台启动 agent（隐藏窗口，日志写入 agent_runtime.log）。",
    "等待 TCP 8898 就绪。",
    "当前终端进入 WASD 遥控。",
    "退出时自动关闭 agent。",
]:
    bullet(doc, item)
body(doc, "只检查环境、不进入遥控：", bold=True)
code(doc, "python start_fishbot.py --check")
body(doc, "若检测到已有 agent 在 8898 上运行，会直接复用；退出遥控时不会关闭被复用的 agent。",
     color="595959")

heading(doc, "推荐启动顺序", 2)
body(doc, "先运行 start_fishbot.py（或先单独启动 agent），再给小车断电重启，"
          "等 3–5 秒完成握手后开始遥控。这样可以保证主控板建立完整的订阅会话。")

heading(doc, "4. 分步启动（两个终端）", 1)
body(doc, "终端 1（agent）：", bold=True)
code(doc, 'cd "C:\\Users\\kw188\\Desktop\\ros小车\\xrce_windows_agent"\n'
          "python -m xrce.agent --host 0.0.0.0 --port 8888")
body(doc, "终端 2（WASD）：", bold=True)
code(doc, 'cd "C:\\Users\\kw188\\Desktop\\ros小车\\xrce_windows_agent"\n'
          "python tools\\teleop_wasd.py")
body(doc, "关闭顺序：终端 2 按 X 退出遥控，终端 1 按 Ctrl+C 关闭 agent。")

heading(doc, "5. 遥控按键", 1)
table(doc, ["按键", "动作", "默认值"], [
    ("W", "前进", "0.05 m/s"),
    ("S", "后退", "0.05 m/s"),
    ("A", "左转", "0.3 rad/s"),
    ("D", "右转", "0.3 rad/s"),
    ("空格", "立即停车", "—"),
    ("X / Ctrl+C", "停车并退出", "—"),
], widths=[1.4, 2.4, 2.4])
bullet(doc, "遥控直接轮询按键的物理状态，不依赖 Windows 键盘连发。")
bullet(doc, "按住方向键时每 0.1 秒刷新一条 0.3 秒的运动指令，连续运动无断档。")
bullet(doc, "松开按键立即发送停车指令。")
bullet(doc, "agent 内部只有一个 20 Hz 发布线程；按键只更新目标速度，不会堆积线程。")

heading(doc, "6. 如何关闭", 1)
heading(doc, "6.1 一键启动方式", 2)
body(doc, "在 WASD 界面按 X 或 Ctrl+C，agent 会在后台自动关闭。")
heading(doc, "6.2 单独启动的 agent", 2)
body(doc, "在 agent 所在终端按 Ctrl+C，这是最干净的退出方式。")
heading(doc, "6.3 终端已关闭 / 进程残留", 2)
body(doc, "PowerShell 按端口关闭：")
code(doc, "Get-NetUDPEndpoint -LocalPort 8888 | "
          "ForEach-Object { Stop-Process -Id $_.OwningProcess -Force }")
body(doc, "或按命令行匹配：")
code(doc, "Get-CimInstance Win32_Process -Filter \"Name='python.exe'\" |\n"
          "  Where-Object { $_.CommandLine -like '*xrce.agent*' } |\n"
          "  ForEach-Object { Stop-Process -Id $_.ProcessId -Force }")
body(doc, "确认已经关闭：")
code(doc, "Get-NetUDPEndpoint -LocalPort 8888")
body(doc, "没有输出即表示 8888 已释放。")
heading(doc, "6.4 切换回 WSL / Docker 方案", 2)
body(doc, "1. 先关闭 Windows 上的 Python agent，避免 8888 冲突。")
body(doc, "2. 在 WSL 中执行：docker start micro_ros_agent")

heading(doc, "7. 常见问题", 1)
table(doc, ["现象", "处理"], [
    ("agent 启动失败", "查看 agent_runtime.log；确认 8888 未被 Docker/WSL agent 占用"),
    ("8898 被占用", "换控制口：--control-port 8897，并同步修改 teleop_wasd.py 的 PORT"),
    ("车不动，但终端返回 OK", "确认上位机 IP 正确；先启动 agent 再重启小车；查看日志是否有 WRITE_DATA"),
    ("未收到 /odom", "主控板未连上或会话不完整，先断电重启小车"),
    ("轮子一顿一顿", "确认使用本版代码（单线程 20 Hz 发布），不要同时开多个 agent"),
    ("中文乱码", "使用 Windows Terminal；脚本已自动切换 UTF-8 代码页"),
    ("会话反复重建", "确认是本项目最新版（会话内 INFO 应答的 impl 字节必须为 01）"),
], widths=[2.1, 4.2])

heading(doc, "8. 常用命令速查", 1)
table(doc, ["目的", "命令"], [
    ("一键启动", "python start_fishbot.py"),
    ("只检查", "python start_fishbot.py --check"),
    ("单独启动 agent", "python -m xrce.agent --host 0.0.0.0 --port 8888"),
    ("WASD 遥控", "python tools\\teleop_wasd.py"),
    ("单条前进", "python tools\\send_cmd.py 0.05 0 0.4"),
    ("单条左转", "python tools\\send_cmd.py 0 0.3 0.4"),
    ("单条后退", "python tools\\send_cmd.py -0.05 0 0.4"),
    ("关闭 agent", "Get-NetUDPEndpoint -LocalPort 8888 | ForEach-Object { Stop-Process -Id $_.OwningProcess -Force }"),
    ("查看日志", "Get-Content agent_runtime.log -Tail 50"),
    ("运行单测", "python -m unittest discover -s tests -v"),
], widths=[2.1, 4.2])

heading(doc, "9. 文件位置", 1)
code(doc, "C:\\Users\\kw188\\Desktop\\ros小车\\xrce_windows_agent\\\n"
          "├─ start_fishbot.py         一键启动\n"
          "├─ agent_runtime.log        一键启动时的 agent 日志\n"
          "├─ xrce\\agent.py            XRCE-DDS agent\n"
          "├─ tools\\teleop_wasd.py     WASD 遥控\n"
          "├─ tools\\send_cmd.py        单条运动指令\n"
          "├─ tests\\test_codec.py      编解码单测\n"
          "└─ docs\\启动关闭手册.md      本手册")

out = "FishBot-Windows原生控制-启动关闭手册.docx"
doc.save(out)
print("saved:", out)
